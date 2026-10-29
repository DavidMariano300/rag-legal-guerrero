"""Mecanismo extensible de generación de documentos.

Este módulo define el "motor" (registro de plantillas + generación de .docx), no el contenido
legal real de ningún documento. La única plantilla registrada por ahora es una de PRUEBA que
demuestra el mecanismo end-to-end; agregar un caso real (ej. demanda alimenticia) implica
únicamente añadir una entrada nueva a TEMPLATES con su propio texto revisado por un abogado,
sin tocar el resto del sistema.
"""

import io

from docx import Document as DocxDocument
from fastapi import HTTPException
from jinja2 import Template

CampoTipo = str  # "texto" | "numero" | "fecha"


class CampoPlantilla:
    def __init__(self, nombre: str, etiqueta: str, tipo: CampoTipo = "texto", requerido: bool = True):
        self.nombre = nombre
        self.etiqueta = etiqueta
        self.tipo = tipo
        self.requerido = requerido

    def to_dict(self) -> dict:
        return {"nombre": self.nombre, "etiqueta": self.etiqueta, "tipo": self.tipo, "requerido": self.requerido}


class PlantillaDocumento:
    def __init__(
        self,
        id: str,
        nombre: str,
        descripcion: str,
        campos: list[CampoPlantilla],
        parrafos: list[str],
        palabras_clave: list[str],
    ):
        self.id = id
        self.nombre = nombre
        self.descripcion = descripcion
        self.campos = campos
        self.parrafos = parrafos  # cada uno puede usar {{ nombre_campo }} (sintaxis Jinja2)
        self.palabras_clave = palabras_clave  # usadas por sugerir_plantilla(); NO se exponen al frontend

    def to_dict(self) -> dict:
        return {
            "id": self.id,
            "nombre": self.nombre,
            "descripcion": self.descripcion,
            "campos": [c.to_dict() for c in self.campos],
        }


TEMPLATES: dict[str, PlantillaDocumento] = {
    "plantilla_prueba": PlantillaDocumento(
        id="plantilla_prueba",
        nombre="Plantilla de prueba (mecanismo, NO es un documento legal real)",
        descripcion=(
            "Documento ficticio usado solo para validar el mecanismo de generación. "
            "No debe usarse en ningún trámite real."
        ),
        campos=[
            CampoPlantilla("nombre_actor", "Nombre del promovente", "texto", True),
            CampoPlantilla("nombre_contraparte", "Nombre de la contraparte", "texto", True),
            CampoPlantilla("fecha", "Fecha", "fecha", True),
            CampoPlantilla("asunto", "Asunto", "texto", False),
        ],
        # Palabras de ejemplo para probar el mecanismo de sugerencia (Sección "sugerir_plantilla").
        # Al agregar una plantilla real, sus palabras clave deben reflejar el tema real del documento.
        # Frases específicas, no palabras sueltas genéricas ("demanda", "indemnización" hubieran
        # coincidido con casi cualquier consulta legal y hubieran repetido el problema original.
        palabras_clave=["pensión alimenticia", "pension alimenticia", "demanda alimenticia"],
        parrafos=[
            "DOCUMENTO DE PRUEBA — NO ES UN ESCRITO LEGAL REAL",
            "Fecha: {{ fecha }}",
            "Promovente: {{ nombre_actor }}",
            "Contraparte: {{ nombre_contraparte }}",
            "Asunto: {{ asunto }}",
            (
                "Este documento fue generado automáticamente por el mecanismo de plantillas de "
                "RAG Legal Guerrero como prueba técnica. Antes de agregar plantillas con contenido "
                "legal real, deben ser redactadas y revisadas por un abogado."
            ),
        ],
    ),
}


def list_templates() -> list[dict]:
    return [t.to_dict() for t in TEMPLATES.values()]


def sugerir_plantilla(pregunta: str, respuesta: str) -> dict | None:
    """Sugiere una plantilla SOLO si el tema de la consulta la amerita.

    Heurística simple por palabras clave para el MVP — no es una decisión del LLM.
    Devuelve None en el caso normal (la mayoría de las consultas no ameritan generar
    un documento), y el dict de la plantilla solo cuando hay una coincidencia real.
    """
    texto = f"{pregunta} {respuesta}".lower()
    for template in TEMPLATES.values():
        if any(palabra in texto for palabra in template.palabras_clave):
            return {"id": template.id, "nombre": template.nombre}
    return None


def generate_document(template_id: str, values: dict[str, str]) -> bytes:
    template = TEMPLATES.get(template_id)
    if template is None:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada.")

    faltantes = [c.etiqueta for c in template.campos if c.requerido and not values.get(c.nombre)]
    if faltantes:
        raise HTTPException(status_code=400, detail=f"Faltan campos requeridos: {', '.join(faltantes)}")

    doc = DocxDocument()
    doc.add_heading(template.nombre, level=1)

    for parrafo_template in template.parrafos:
        texto = Template(parrafo_template).render(**values)
        doc.add_paragraph(texto)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
